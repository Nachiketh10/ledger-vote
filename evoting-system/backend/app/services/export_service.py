"""
Generates downloadable PDF and Excel exports of an election's results.
Kept as pure functions (dict in, bytes out) — no Flask/DB imports here,
so this is easy to unit test and reuse (e.g. a future "email me the
results" feature could call the same functions).
"""

import io
from datetime import datetime, timezone


def generate_results_pdf(results: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleCustom", parent=styles["Title"], fontSize=18)
    meta_style = ParagraphStyle("Meta", parent=styles["Normal"], textColor=colors.grey)

    elements = [
        Paragraph(results["title"], title_style),
        Paragraph("Official Election Results", meta_style),
        Spacer(1, 6),
        Paragraph(
            f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} · "
            f"Blockchain integrity: {'VERIFIED' if results['blockchain_verified'] else 'FAILED — DO NOT TRUST THESE RESULTS'}",
            meta_style,
        ),
        Spacer(1, 20),
    ]

    summary_data = [
        ["Total votes cast", str(results["total_votes"])],
        ["Eligible voters", str(results["eligible_voter_count"])],
        ["Participation", f"{results['participation_percentage']}%"],
    ]
    summary_table = Table(summary_data, colWidths=[2.5 * inch, 2 * inch])
    summary_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ]
        )
    )
    elements += [summary_table, Spacer(1, 24)]

    if results.get("winner"):
        elements += [
            Paragraph(f"<b>Winner: {results['winner']['name']}</b> "
                      f"({results['winner']['votes']} votes, {results['winner']['percentage']}%)",
                      styles["Heading3"]),
            Spacer(1, 12),
        ]

    candidate_rows = [["Candidate", "Party", "Votes", "Percentage"]]
    for c in results["candidates"]:
        candidate_rows.append([c["name"], c.get("party") or "—", str(c["votes"]), f"{c['percentage']}%"])

    candidate_table = Table(candidate_rows, colWidths=[2.2 * inch, 1.8 * inch, 1 * inch, 1.2 * inch])
    candidate_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24406B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("TOPPADDING", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F4EE")]),
            ]
        )
    )
    elements.append(candidate_table)

    doc.build(elements)
    return buffer.getvalue()


def generate_results_excel(results: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    header_fill = PatternFill(start_color="24406B", end_color="24406B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)

    ws.append([results["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Blockchain integrity: {'VERIFIED' if results['blockchain_verified'] else 'FAILED'}"])
    ws.append([])

    ws.append(["Total votes cast", results["total_votes"]])
    ws.append(["Eligible voters", results["eligible_voter_count"]])
    ws.append(["Participation %", results["participation_percentage"]])
    ws.append([])

    if results.get("winner"):
        ws.append(["Winner", results["winner"]["name"], f"{results['winner']['votes']} votes"])
        ws["A8"].font = bold
        ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(["Candidate", "Party", "Votes", "Percentage"])
    for cell in ws[header_row_idx]:
        cell.fill = header_fill
        cell.font = header_font

    for c in results["candidates"]:
        ws.append([c["name"], c.get("party") or "—", c["votes"], c["percentage"]])

    for col, width in zip("ABCD", (28, 20, 10, 12)):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
